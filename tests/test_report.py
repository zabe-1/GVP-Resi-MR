"""End-to-end workbook generation from synthetic results (no network)."""

from openpyxl import load_workbook

from cre_market.geo import BlockGroup
from cre_market.pipeline import AddressResult, MetricResult, RadiusResult
from cre_market.report import write_workbook

RADII = (30.0, 15.0, 5.0)


def _result():
    res = AddressResult(address="123 Main St, Houston, TX",
                        matched_address="123 MAIN ST, HOUSTON, TX, 77002",
                        lat=29.76, lon=-95.36, acs_year=2023)
    for r in RADII:
        rr = RadiusResult(radius_miles=r, n_block_groups=int(r * 10),
                          pct_bg_suppressed=3.5, n_bg_high_moe=4,
                          counties=["48201"],
                          costar_record_counts={"rent": 12})
        rr.metrics["population_growth"] = MetricResult(
            value=8.2, unit="%", source="Census ACS 5-year",
            year_range="ACS5 2018 vs 2023")
        rr.metrics["rent_growth"] = MetricResult(
            value=21.0, unit="%", source="CoStar export (rents.csv)",
            year_range="2019-2024", note="12 records in radius")
        rr.metrics["job_growth"] = MetricResult(
            value=None, source="BLS QCEW", note="no data")
        res.radii[r] = rr
    res.block_groups_current = [BlockGroup(
        geoid="482010100001", state="48", county="201", tract="010000",
        blkgrp="1", lat=29.76, lon=-95.36, distance_miles=1.23)]
    res.county_detail = [{"radius_miles": 30.0, "metric": "QCEW jobs",
                          "county_fips": "48201", "county_name": "Harris",
                          "weight": 1.0, "growth_pct": 5.5}]
    return res


def test_workbook_structure(tmp_path):
    path = str(tmp_path / "out.xlsx")
    write_workbook(path, [_result()], RADII, lookback=5)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Coverage & Quality", "Block Groups",
                             "Counties", "Methodology"]
    ws = wb["Summary"]
    assert ws.cell(1, 1).value == "Address"
    # 9 metrics x (3 radii + source col) + address col
    assert ws.max_column == 1 + 9 * 4
    assert ws.cell(3, 1).value == "123 MAIN ST, HOUSTON, TX, 77002"
    assert ws.cell(3, 2).value == 8.2               # population growth, 30mi
    src = ws.cell(3, 5).value                       # population source column
    assert "Census ACS" in src and "2018 vs 2023" in src


def test_workbook_marks_missing_and_costar(tmp_path):
    path = str(tmp_path / "out.xlsx")
    write_workbook(path, [_result()], RADII, lookback=5)
    ws = load_workbook(path)["Summary"]
    header_starts = {}
    col = 2
    for _ in range(9):
        label = ws.cell(1, col).value
        header_starts[label] = col
        col += 4
    jg = header_starts["Job Growth (5yr, %)"]
    assert ws.cell(3, jg).value is None            # unreliable -> blank cell
    assert ws.cell(3, jg).comment is not None      # reason kept as comment
    rg = header_starts["Rental Rate Growth (5yr, %)"]
    assert ws.cell(3, rg).value == 21.0
    assert "CoStar" in ws.cell(3, rg + 3).value      # source column labels it


def test_coverage_sheet_reports_counts(tmp_path):
    path = str(tmp_path / "out.xlsx")
    write_workbook(path, [_result()], RADII, lookback=5)
    ws = load_workbook(path)["Coverage & Quality"]
    assert ws.cell(2, 2).value == 30                 # radius
    assert ws.cell(2, 3).value == 300                # block groups
    assert ws.cell(2, 8).value == 12                 # CoStar rent records


def test_error_address_row(tmp_path):
    bad = AddressResult(address="nowhere", error="geocoding failed: no match")
    path = str(tmp_path / "out.xlsx")
    write_workbook(path, [bad], RADII, lookback=5)
    ws = load_workbook(path)["Summary"]
    assert "ERROR" in ws.cell(3, 2).value
