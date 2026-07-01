"""Excel workbook output.

Layout decision (spec asked for a recommendation): ONE ROW PER ADDRESS on a
single Summary sheet, columns grouped metric -> radius.  Side-by-side rows
make cross-property comparison trivial, which is the point of a market study;
per-address detail that doesn't fit a row goes to dedicated debug sheets.

Sheets:
  Summary            one row per address; per metric: 30/15/5-mile values +
                     source & year range
  Coverage & Quality per address x radius: block-group counts, % suppressed,
                     high-MOE counts, counties touched, CoStar record counts
  Block Groups       every block group included, with distance from the point
  Counties           per-county values behind every county-level metric
  Methodology        sources, year ranges, and caveats in plain English
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .pipeline import METRIC_LABELS, METRIC_ORDER, AddressResult

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
SUBHEADER_FILL = PatternFill("solid", fgColor="D6E4F0")
COSTAR_FILL = PatternFill("solid", fgColor="FFF2CC")   # CoStar-sourced cells
FLAG_FILL = PatternFill("solid", fgColor="FCE4E4")     # missing / failed cells
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Border(*[Side(style="thin", color="BBBBBB")] * 4)


def write_workbook(path: str, results: list[AddressResult],
                   radii: tuple[float, ...], lookback: int) -> None:
    radii = tuple(sorted(radii, reverse=True))
    wb = Workbook()
    _summary_sheet(wb.active, results, radii)
    _coverage_sheet(wb.create_sheet("Coverage & Quality"), results, radii)
    _block_groups_sheet(wb.create_sheet("Block Groups"), results, radii)
    _counties_sheet(wb.create_sheet("Counties"), results)
    _methodology_sheet(wb.create_sheet("Methodology"), results, lookback)
    wb.save(path)


def _summary_sheet(ws, results, radii):
    ws.title = "Summary"
    # Row 1: metric group headers (merged); Row 2: radius / source labels.
    ws.cell(1, 1, "Address").fill = HEADER_FILL
    ws.cell(1, 1).font = HEADER_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    col = 2
    metric_cols: dict[str, int] = {}
    for key in METRIC_ORDER:
        span = len(radii) + 1  # radii + Source column
        ws.merge_cells(start_row=1, start_column=col, end_row=1,
                       end_column=col + span - 1)
        c = ws.cell(1, col, METRIC_LABELS[key])
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        metric_cols[key] = col
        for i, r in enumerate(radii):
            sc = ws.cell(2, col + i, f"{r:g} mi")
            sc.fill = SUBHEADER_FILL
            sc.font = Font(bold=True)
        sc = ws.cell(2, col + len(radii), "Source / Years")
        sc.fill = SUBHEADER_FILL
        sc.font = Font(bold=True)
        col += span

    for row_i, res in enumerate(results, start=3):
        ws.cell(row_i, 1, res.matched_address or res.address)
        if res.error:
            c = ws.cell(row_i, 2, f"ERROR: {res.error}")
            c.fill = FLAG_FILL
            continue
        for key in METRIC_ORDER:
            base = metric_cols[key]
            source_texts = set()
            for i, r in enumerate(radii):
                mr = res.radii.get(r) and res.radii[r].metrics.get(key)
                cell = ws.cell(row_i, base + i)
                if mr is None or mr.value is None:
                    cell.value = "n/a"
                    cell.fill = FLAG_FILL
                    if mr and mr.note:
                        cell.comment = _comment(mr.note)
                else:
                    cell.value = round(mr.value, 1)
                    cell.number_format = "#,##0.0" if mr.unit == "%" else "#,##0"
                    if "CoStar" in mr.source:
                        cell.fill = COSTAR_FILL
                    if mr.note:
                        cell.comment = _comment(mr.note)
                if mr:
                    label = mr.source + (f" [{mr.year_range}]" if mr.year_range else "")
                    source_texts.add(label)
            sc = ws.cell(row_i, base + len(radii), " | ".join(sorted(source_texts)))
            sc.alignment = Alignment(wrap_text=True)
            sc.font = Font(size=8)

    for c in range(1, col):
        ws.column_dimensions[get_column_letter(c)].width = 13 if c > 1 else 40
    for key in METRIC_ORDER:
        ws.column_dimensions[get_column_letter(metric_cols[key] + len(radii))].width = 32
    ws.freeze_panes = "B3"


def _comment(text: str):
    from openpyxl.comments import Comment
    return Comment(text[:700], "cre-market")


def _coverage_sheet(ws, results, radii):
    headers = ["Address", "Radius (mi)", "Block groups included",
               "% BGs with suppressed ACS data", "BGs with high-MOE values",
               "Counties touched", "County FIPS list",
               "CoStar rent records", "CoStar home-sale records",
               "CoStar new-home-sale records"]
    _header_row(ws, headers)
    row = 2
    for res in results:
        for r in sorted(radii, reverse=True):
            rr = res.radii.get(r)
            if rr is None:
                continue
            ws.cell(row, 1, res.matched_address or res.address)
            ws.cell(row, 2, r)
            ws.cell(row, 3, rr.n_block_groups)
            ws.cell(row, 4, round(rr.pct_bg_suppressed, 1))
            ws.cell(row, 5, rr.n_bg_high_moe)
            ws.cell(row, 6, len(rr.counties))
            ws.cell(row, 7, ", ".join(rr.counties))
            ws.cell(row, 8, rr.costar_record_counts.get("rent"))
            ws.cell(row, 9, rr.costar_record_counts.get("home_sales"))
            ws.cell(row, 10, rr.costar_record_counts.get("new_home_sales"))
            row += 1
    _autosize(ws, headers)


def _block_groups_sheet(ws, results, radii):
    headers = ["Address", "ACS vintage", "Block group GEOID", "County FIPS",
               "Distance (mi)", "Within radii"]
    _header_row(ws, headers)
    row = 2
    for res in results:
        for label, bgs in (("current", res.block_groups_current),
                           ("prior", res.block_groups_prior)):
            for bg in bgs:
                within = [f"{r:g}" for r in sorted(radii)
                          if bg.distance_miles <= r]
                ws.cell(row, 1, res.matched_address or res.address)
                ws.cell(row, 2, f"{res.acs_year} ({label})" if label == "current"
                        else f"prior vintage")
                ws.cell(row, 3, bg.geoid)
                ws.cell(row, 4, bg.county_fips)
                ws.cell(row, 5, round(bg.distance_miles, 2))
                ws.cell(row, 6, ", ".join(within) + " mi")
                row += 1
    _autosize(ws, headers)


def _counties_sheet(ws, results):
    all_keys: list[str] = []
    for res in results:
        for rec in res.county_detail:
            for k in rec:
                if k not in all_keys:
                    all_keys.append(k)
    headers = ["Address"] + all_keys
    _header_row(ws, headers)
    row = 2
    for res in results:
        for rec in res.county_detail:
            ws.cell(row, 1, res.matched_address or res.address)
            for j, k in enumerate(all_keys, start=2):
                v = rec.get(k)
                ws.cell(row, j, round(v, 2) if isinstance(v, float) else v)
            row += 1
    _autosize(ws, headers)


def _methodology_sheet(ws, results, lookback):
    lines = [
        "METHODOLOGY & SOURCES",
        "",
        "Radius selection: Census block groups whose centroid (internal point) falls inside "
        "each radius (centroid-in-radius method, not boundary overlap). Block groups are "
        "selected independently for each ACS vintage because 2010-based and 2020-based "
        "boundaries differ; growth compares aggregates inside the same physical circle.",
        "",
        f"Lookback: {lookback} years. Every figure's exact year range appears in the "
        "Summary sheet's Source / Years column.",
        "",
        "Metric sources (government defaults):",
        " 1. Population Growth — Census ACS 5-year B01003, sum of block-group populations.",
        " 2. HH Income Growth — Census ACS 5-year B19013; household-weighted average of "
        "block-group median incomes (approximation of area median).",
        " 3. Household Formations — Census ACS 5-year B25002_002 (occupied housing units), "
        "net change over the lookback.",
        " 4. Employment Rate — Census ACS 5-year B23025: employed / civilian labor force 16+ "
        "(current vintage level).",
        " 5. Educational Attainment — Census ACS 5-year B15003: % of pop 25+ with bachelor's "
        "degree or higher (current vintage level).",
        " 6. Job Growth — BLS QCEW annual average employment, county totals summed across "
        "counties touched by the radius.",
        " 7. New Home Sales Growth — Census Building Permits Survey, 1-unit permits (PROXY "
        "for new home sales, not closings), summed across counties. Overridden by CoStar "
        "export when provided.",
        " 8. Rental Rate Growth — HUD Fair Market Rents, 2BR, county level; county growth "
        "weighted by radius population share. Overridden by CoStar export when provided.",
        " 9. Existing Home Sales Growth — FHFA all-transactions county HPI (PRICE "
        "APPRECIATION proxy; no free government source publishes county sales volume). "
        "Overridden by CoStar export when provided.",
        "",
        "CoStar-sourced cells are shaded yellow in the Summary sheet, and the Source column "
        "names the export file. CoStar records are re-filtered to each radius using the "
        "geocoded point; per-radius record counts are in Coverage & Quality.",
        "",
        "Reliability: suppressed ACS values (sentinel codes) are excluded from aggregates "
        "but counted and reported; estimates with CV > 30% are flagged as high-MOE. "
        "Red cells in the Summary sheet mean no value could be computed — hover the cell "
        "comment for the reason.",
        "",
        "ACS caveat: 5-year estimates are pooled periods (vintage 2023 = 2019-2023), so a "
        f"{lookback}-year vintage comparison reflects two pooled windows, not two point-in-"
        "time years.",
    ]
    for i, line in enumerate(lines, start=1):
        ws.cell(i, 1, line)
    ws.column_dimensions["A"].width = 120
    for i in (1,):
        ws.cell(i, 1).font = Font(bold=True, size=12)


def _header_row(ws, headers):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(1, j, h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = "A2"


def _autosize(ws, headers):
    for j, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(14, min(40, len(h) + 4))
